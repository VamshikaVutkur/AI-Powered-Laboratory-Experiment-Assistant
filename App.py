import streamlit as st
# Load CSS
with open("style.css", "r", encoding="utf-8") as f:
    st.markdown(
        f"<style>{f.read()}</style>",
        unsafe_allow_html=True
    )
    
import os
import shutil
import streamlit as st

from pypdf import PdfReader
from docx import Document

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document as LCDocument
from langchain_core.prompts import ChatPromptTemplate


# ---------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------

st.set_page_config(
    page_title="AI Laboratory Experiment Assistant",
    page_icon="🧪",
    layout="wide"
)


# ---------------------------------------------------------
# TITLE
# ---------------------------------------------------------

st.title("🧪 AI-Powered Laboratory Experiment Assistant")

st.write(
    "Upload your laboratory manuals and ask questions about "
    "experiments, procedures, results, precautions, and viva questions."
)


# ---------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------

CHROMA_PATH = "chroma_db"

LLM_MODEL = "llama3.2"
EMBEDDING_MODEL = "nomic-embed-text"


# ---------------------------------------------------------
# INITIALIZE OLLAMA
# ---------------------------------------------------------

@st.cache_resource
def get_models():

    embeddings = OllamaEmbeddings(
        model=EMBEDDING_MODEL
    )

    llm = ChatOllama(
        model=LLM_MODEL,
        temperature=0.2
    )

    return embeddings, llm


# ---------------------------------------------------------
# READ PDF
# ---------------------------------------------------------

def read_pdf(uploaded_file):

    reader = PdfReader(uploaded_file)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ---------------------------------------------------------
# READ DOCX
# ---------------------------------------------------------

def read_docx(uploaded_file):

    document = Document(uploaded_file)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text += paragraph.text + "\n"

    return text


# ---------------------------------------------------------
# PROCESS FILE
# ---------------------------------------------------------

def process_file(uploaded_file):

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):

        text = read_pdf(uploaded_file)

    elif file_name.endswith(".docx"):

        text = read_docx(uploaded_file)

    else:

        return []

    if not text.strip():

        return []

    document = LCDocument(
        page_content=text,
        metadata={
            "source": uploaded_file.name
        }
    )

    return [document]


# ---------------------------------------------------------
# SPLIT DOCUMENTS
# ---------------------------------------------------------

def split_documents(documents):

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200
    )

    chunks = text_splitter.split_documents(documents)

    return chunks


# ---------------------------------------------------------
# CREATE VECTOR DATABASE
# ---------------------------------------------------------

def create_vector_database(chunks, embeddings):

    vectorstore = Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        persist_directory=CHROMA_PATH
    )

    return vectorstore


# ---------------------------------------------------------
# ANSWER QUESTION
# ---------------------------------------------------------

def answer_question(question, vectorstore, llm, chat_history=None):

    # Get relevant information from laboratory documents
    relevant_documents = []

    if vectorstore is not None:

        try:

            relevant_documents = vectorstore.similarity_search(
                question,
                k=4
            )

        except Exception:
            relevant_documents = []

    # Create context from laboratory documents
    context = ""

    if relevant_documents:

        context = "\n\n".join(
            document.page_content
            for document in relevant_documents
        )

    # Create conversation history
    history_text = ""

    if chat_history:

        for message in chat_history[-6:]:

            history_text += (
                f"{message['role']}: "
                f"{message['content']}\n"
            )

    # Prompt
    prompt = ChatPromptTemplate.from_template(
        """
You are an intelligent AI-powered Laboratory Experiment Assistant
and general-purpose student chatbot.

You can answer BOTH:
1. Questions related to the uploaded laboratory material.
2. General questions asked by the student.

IMPORTANT RULES:

- If the student's question is related to a laboratory experiment,
  use the laboratory material provided below whenever relevant.
- If the question is NOT related to the laboratory material,
  answer it normally using your general knowledge.
- Do NOT say that you can only answer questions from the PDF.
- Explain answers clearly and in a student-friendly way.
- Use headings, bullet points, and numbered steps when useful.
- If the student asks for code, provide working code when appropriate.
- If the student asks a simple question, give a simple answer.
- If the student asks for a detailed explanation, give a detailed answer.
- Never pretend that information from the laboratory manual exists
  if it is not actually present.

LABORATORY MATERIAL:
-------------------------
{context}
-------------------------

PREVIOUS CONVERSATION:
-------------------------
{history}
-------------------------

STUDENT QUESTION:
-------------------------
{question}
-------------------------

ANSWER:
"""
    )

    chain = prompt | llm

    response = chain.invoke(
        {
            "context": context,
            "history": history_text,
            "question": question
        }
    )

    return response.content

# ---------------------------------------------------------
# SIDEBAR
# ---------------------------------------------------------

with st.sidebar:

    st.header("📚 Laboratory Materials")

    uploaded_files = st.file_uploader(
        "Upload laboratory manuals",
        type=["pdf", "docx"],
        accept_multiple_files=True
    )

    build_database = st.button(
        "🔨 Build Knowledge Base",
        use_container_width=True
    )


# ---------------------------------------------------------
# BUILD KNOWLEDGE BASE
# ---------------------------------------------------------

if build_database:

    if not uploaded_files:

        st.warning(
            "Please upload at least one PDF or DOCX file."
        )

    else:

        with st.spinner(
            "Processing laboratory documents..."
        ):

            try:

                embeddings, llm = get_models()

                all_documents = []

                for uploaded_file in uploaded_files:

                    documents = process_file(
                        uploaded_file
                    )

                    all_documents.extend(
                        documents
                    )

                if not all_documents:

                    st.error(
                        "Could not extract text from the uploaded files."
                    )

                else:

                    chunks = split_documents(
                        all_documents
                    )

                    create_vector_database(
                        chunks,
                        embeddings
                    )

                    st.session_state["database_ready"] = True

                    st.success(
                        f"Knowledge base created successfully!"
                    )

                    st.info(
                        f"Documents: {len(all_documents)} | "
                        f"Text chunks: {len(chunks)}"
                    )

            except Exception as e:

                st.error(
                    "An error occurred while creating the knowledge base."
                )

                st.exception(e)


# ---------------------------------------------------------
# QUESTION SECTION
# ---------------------------------------------------------

# ---------------------------------------------------------
# CHATBOT
# ---------------------------------------------------------

st.header("💬 AI Laboratory Chatbot")

# Create chat history
if "messages" not in st.session_state:

    st.session_state.messages = []


# Display previous messages
for message in st.session_state.messages:

    with st.chat_message(message["role"]):

        st.markdown(message["content"])


# Chat input
question = st.chat_input(
    "Ask me anything about your laboratory or any other topic..."
)


# Process question
if question:

    # Display user question
    st.session_state.messages.append(
        {
            "role": "user",
            "content": question
        }
    )

    with st.chat_message("user"):

        st.markdown(question)


    # Get AI response
    with st.chat_message("assistant"):

        with st.spinner("Thinking..."):

            try:

                embeddings, llm = get_models()

                # Load Chroma database if available
                if os.path.exists(CHROMA_PATH):

                    vectorstore = Chroma(
                        persist_directory=CHROMA_PATH,
                        embedding_function=embeddings
                    )

                else:

                    vectorstore = None


                answer = answer_question(
                    question,
                    vectorstore,
                    llm,
                    st.session_state.messages
                )

                st.markdown(answer)


                # Save AI response
                st.session_state.messages.append(
                    {
                        "role": "assistant",
                        "content": answer
                    }
                )


            except Exception as e:

                st.error(
                    "Something went wrong while generating the answer."
                )

                st.exception(e)

# ---------------------------------------------------------
# QUICK QUESTIONS
# ---------------------------------------------------------

st.header("⚡ Quick Questions")

col1, col2, col3 = st.columns(3)


# ---------------------------------------------------------
# EXPLAIN EXPERIMENT
# ---------------------------------------------------------

with col1:

    if st.button(
        "📖 Explain Experiment",
        use_container_width=True
    ):

        if not os.path.exists(CHROMA_PATH):

            st.warning(
                "Please upload your laboratory manual "
                "and click 'Build Knowledge Base' first."
            )

        else:

            with st.spinner("Explaining the experiment..."):

                try:

                    embeddings, llm = get_models()

                    vectorstore = Chroma(
                        persist_directory=CHROMA_PATH,
                        embedding_function=embeddings
                    )

                    answer = answer_question(
                        "Explain the experiment in detail. "
                        "Include the aim, theory, procedure, "
                        "requirements, and expected result.",
                        vectorstore,
                        llm
                    )

                    st.subheader("📖 Experiment Explanation")

                    st.write(answer)

                except Exception as e:

                    st.error("Error generating explanation.")

                    st.exception(e)

# ---------------------------------------------------------
# VIVA QUESTIONS
# ---------------------------------------------------------

with col2:
    if st.button(
        "🎓 Viva Questions",
        use_container_width=True
    ):

        if not os.path.exists(CHROMA_PATH):

            st.warning(
                "Please upload your laboratory manual "
                "and click 'Build Knowledge Base' first."
            )

        else:

            with st.spinner("Generating AI viva questions..."):

                try:

                    embeddings, llm = get_models()

                    vectorstore = Chroma(
                        persist_directory=CHROMA_PATH,
                        embedding_function=embeddings
                    )

                    # First get the experiment information from the manual
                    experiment_info = answer_question(
                        "Explain the experiment briefly. "
                        "Include the experiment name, aim, theory, "
                        "principle, important concepts, procedure, "
                        "equipment, and expected result.",
                        vectorstore,
                        llm
                    )

                    # Ask the AI to CREATE its own viva questions
                    viva_prompt = f"""
You are a laboratory viva examiner.

Based on the following information about the experiment:

{experiment_info}

Generate 10 important viva questions and answers.

IMPORTANT INSTRUCTIONS:
- The questions do NOT have to appear in the laboratory manual.
- Create new questions using your own knowledge about the experiment.
- Include conceptual, theoretical, practical, and application-based questions.
- Ask questions that a college professor might actually ask in a viva.
- Give a clear and correct answer after every question.
- Do not say that the questions are not in the manual.
- Do not search for the exact questions in the manual.
- Make the questions relevant to this specific experiment.

Format:

Q1. Question
Answer: ...

Q2. Question
Answer: ...

Continue until Q10.
"""

                    response = llm.invoke(viva_prompt)

                    st.subheader("🎓 AI Generated Viva Questions")

                    st.write(response.content)

                except Exception as e:

                    st.error("Error generating viva questions.")

                    st.exception(e)

# ---------------------------------------------------------
# PRECAUTIONS
# ---------------------------------------------------------


# ---------------------------------------------------------
# SAFETY & PRECAUTIONS
# ---------------------------------------------------------

st.header("⚠️ Safety & Precautions")

col1, col2 = st.columns(2)


# =========================================================
# EXPERIMENT-SPECIFIC PRECAUTIONS
# =========================================================

with col1:

    if st.button(
        "🧪 Experiment Precautions",
        use_container_width=True
    ):

        if not os.path.exists(CHROMA_PATH):

            st.warning(
                "Please upload your laboratory manual "
                "and click 'Build Knowledge Base' first."
            )

        else:

            with st.spinner(
                "Finding experiment-specific precautions..."
            ):

                try:

                    embeddings, llm = get_models()

                    vectorstore = Chroma(
                        persist_directory=CHROMA_PATH,
                        embedding_function=embeddings
                    )

                    precaution_question = """
                    Identify the type of laboratory and the experiment
                    from the uploaded laboratory manual.

                    Then provide ONLY precautions that are relevant
                    to that laboratory and experiment.

                    IMPORTANT:

                    - Do NOT give generic safety precautions.
                    - Do NOT assume this is a chemistry laboratory.
                    - Do NOT mention fumes, flames, chemicals,
                      chemical spills, fire hazards, or chemical
                      handling unless the uploaded laboratory manual
                      actually contains such activities.
                    - If this is a DBMS or computer laboratory,
                      give DBMS/computer-related precautions.
                    - If this is a programming laboratory,
                      give programming/computer-related precautions.
                    - If this is an electronics laboratory,
                      give electronics/electrical precautions.
                    - If this is an electrical laboratory,
                      give electrical safety precautions.
                    - If this is a chemistry laboratory,
                      chemical precautions may be given only when
                      they are relevant to the experiment.
                    - If the manual specifically states precautions,
                      prioritize those precautions.

                    For a DBMS laboratory, consider relevant points
                    such as:
                    - Checking SQL syntax before execution.
                    - Using the correct database and table.
                    - Being careful with DELETE, UPDATE and DROP.
                    - Checking WHERE conditions before modifying data.
                    - Avoiding accidental data loss.
                    - Saving SQL queries and results.
                    - Using database permissions correctly.
                    - Following the experiment procedure.

                    Return the answer in this format:

                    ### 🧪 Experiment-Specific Precautions

                    1. ...
                    2. ...
                    3. ...
                    4. ...

                    ### ⚠️ Things to Avoid

                    1. ...
                    2. ...
                    3. ...

                    Make the answer specific to the experiment.
                    Do not invent hazards that are not relevant.
                    """

                    answer = answer_question(
                        precaution_question,
                        vectorstore,
                        llm,
                        st.session_state.get(
                            "messages",
                            []
                        )
                    )

                    st.subheader(
                        "🧪 Experiment-Specific Precautions"
                    )

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        "Error finding experiment precautions."
                    )

                    st.exception(e)


# =========================================================
# GENERAL LABORATORY SAFETY RULES
# =========================================================

with col2:

    if st.button(
        "🏫 Laboratory Safety Rules",
        use_container_width=True
    ):

        with st.spinner(
            "Preparing laboratory safety rules..."
        ):

            try:

                embeddings, llm = get_models()

                # -------------------------------------------------
                # GET LABORATORY INFORMATION FROM THE PDF
                # -------------------------------------------------

                if os.path.exists(CHROMA_PATH):

                    vectorstore = Chroma(
                        persist_directory=CHROMA_PATH,
                        embedding_function=embeddings
                    )

                    lab_documents = vectorstore.similarity_search(
                        "What type of laboratory is this manual for?",
                        k=5
                    )

                    lab_context = "\n\n".join(
                        doc.page_content
                        for doc in lab_documents
                    )

                else:

                    lab_context = ""

                # -------------------------------------------------
                # SAFETY PROMPT
                # -------------------------------------------------

                safety_prompt = """
                You are an AI Laboratory Safety Assistant.

                Determine the type of laboratory from the
                laboratory material provided below.

                The laboratory could be:

                - DBMS / Computer Science
                - Programming
                - Computer Networks
                - Data Structures
                - Electronics
                - Electrical
                - Physics
                - Chemistry
                - Biology
                - Any other academic laboratory

                IMPORTANT RULES:

                1. Give safety rules appropriate ONLY to the
                   laboratory type.

                2. DO NOT give chemistry laboratory rules to a
                   DBMS, programming, networking, or computer lab.

                3. DO NOT mention:
                   - fumes
                   - flames
                   - chemicals
                   - chemical spills
                   - chemical handling

                   unless the laboratory material actually
                   involves these things.

                4. Do not invent hazards.

                5. Use the laboratory material as the primary
                   source whenever possible.

                6. If the laboratory is a DBMS or computer lab,
                   focus on:
                   - proper computer usage
                   - database/data safety
                   - SQL command safety
                   - protecting data
                   - proper system usage
                   - avoiding unauthorized software
                   - proper shutdown/log out
                   - protecting user accounts and passwords

                7. If the laboratory is an electronics or
                   electrical lab, provide appropriate electrical
                   and equipment precautions.

                8. If the laboratory is a chemistry lab, provide
                   appropriate chemical safety only when relevant.

                9. Keep the rules simple and suitable for
                   college students.

                Organize the answer using:

                ### 🧤 Before Starting

                Give appropriate preparation and safety checks.

                ### 🧪 During the Experiment

                Give safety rules relevant to the laboratory.

                ### ⚠️ Things to Avoid

                Give actions students should avoid.

                ### 🧹 After the Experiment

                Give appropriate cleanup, saving, shutdown,
                and equipment rules.

                ### 🚨 Emergency / Problem Handling

                Give appropriate instructions if something
                goes wrong.

                Do NOT include irrelevant laboratory hazards.

                LABORATORY MATERIAL:
                -------------------------
                {context}
                -------------------------
                """

                prompt = ChatPromptTemplate.from_template(
                    safety_prompt
                )

                chain = prompt | llm

                response = chain.invoke(
                    {
                        "context": lab_context
                    }
                )

                st.subheader(
                    "🏫 Laboratory Safety Rules"
                )

                st.markdown(
                    response.content
                )

            except Exception as e:

                st.error(
                    "Error generating laboratory safety rules."
                )

                st.exception(e)